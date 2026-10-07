"""Оркестратор AI-парсинга страниц.

Обходит структуру: base_dir/название_книги/convert_in_jpg/*.jpg
Обрабатывает каждую страницу через 3-шаговый пайплайн (page_parser.py).
Сохраняет результаты:
  base_dir/название_книги/parsed/*.md          — Markdown с текстом
  base_dir/название_книги/parsed/images/*.jpg  — вырезанные картинки
  base_dir/название_книги/parsed/images_descriptions.json
  base_dir/название_книги/parsed/parsing_report.csv

Поддерживает:
- Возобновление (пропускает уже обработанные файлы по CSV)
- Батчевое сохранение (каждые batch_size файлов)
- Отмену через threading.Event
- Обратный вызов прогресса
"""

from __future__ import annotations

import csv
import json
import logging
import os
import threading
import time
from io import BytesIO
from pathlib import Path
from typing import Callable

from src.config import setup_logging, DATA_DIR
from src.llm_clients import FatalAPIError
from src.page_parser import (
    process_single_page,
    process_page_step1,
    process_page_step2,
    process_page_step3,
)

logger = logging.getLogger(__name__)

# Расширения изображений для поиска
JPG_EXTS = (".jpg", ".jpeg", ".png")


def _find_jpg_files(
    base_dir: Path,
    book_name: str | None = None,
) -> list[tuple[str, Path, Path]]:
    """Найти все JPG в подпапках */convert_in_jpg/.

    Args:
        base_dir:  Корневая папка с книгами.
        book_name: Если указано — искать только в этой папке книги.

    Returns:
        Список кортежей: (book_name, jpg_path, parsed_output_dir)
    """
    results: list[tuple[str, Path, Path]] = []

    if not base_dir.exists():
        return results

    # Определяем список папок для обхода
    if book_name:
        candidate = base_dir / book_name
        if candidate.is_dir():
            book_dirs = [candidate]
        else:
            logger.warning("Папка книги не найдена: %s", candidate)
            return results
    else:
        book_dirs = sorted(base_dir.iterdir())

    for book_dir in book_dirs:
        if not book_dir.is_dir():
            continue

        jpg_src_dir = book_dir / "convert_in_jpg"
        if not jpg_src_dir.is_dir():
            logger.debug("Пропуск %s: нет convert_in_jpg/", book_dir.name)
            continue

        parsed_dir = book_dir / "parsed"

        for jpg_file in sorted(jpg_src_dir.glob("*")):
            if jpg_file.is_file() and jpg_file.suffix.lower() in JPG_EXTS:
                results.append((book_dir.name, jpg_file, parsed_dir))

    return results


def _load_previous_progress(
    parsed_dir: Path,
) -> tuple[list[list], list[dict], set[str]]:
    """Загрузить предыдущий прогресс из CSV и JSON.

    Returns:
        (csv_rows, images_data, processed_filenames)
    """
    csv_rows: list[list] = []
    images_data: list[dict] = []
    processed: set[str] = set()

    csv_path = parsed_dir / "parsing_report.csv"
    if csv_path.exists():
        try:
            with open(csv_path, "r", encoding="utf-8") as f:
                reader = csv.reader(f)
                header = next(reader, None)  # пропускаем заголовок
                for row in reader:
                    if row:
                        csv_rows.append(row)
                        processed.add(row[0])  # filename — первый столбец
            logger.info("Загружен CSV: %d записей", len(csv_rows))
        except Exception as e:
            logger.warning("Ошибка чтения CSV: %s", e)

    json_path = parsed_dir / "images_descriptions.json"
    if json_path.exists():
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                images_data = json.load(f)
            logger.info("Загружен JSON: %d записей", len(images_data))
        except Exception as e:
            logger.warning("Ошибка чтения JSON: %s", e)

    return csv_rows, images_data, processed


def _save_progress(
    parsed_dir: Path,
    csv_rows: list[list],
    images_data: list[dict],
) -> None:
    """Сохранить прогресс в CSV и JSON."""
    parsed_dir.mkdir(parents=True, exist_ok=True)

    # JSON
    json_path = parsed_dir / "images_descriptions.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(images_data, f, ensure_ascii=False, indent=2)

    # CSV
    csv_path = parsed_dir / "parsing_report.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Filename", "Time_sec", "Has_Image",
            "Step1_Prompt", "Step1_Comp",
            "Step2_Prompt", "Step2_Comp",
            "Step3_Prompt", "Step3_Comp",
            "Total_Prompt", "Total_Comp", "Grand_Total",
        ])
        writer.writerows(csv_rows)

    logger.info("Прогресс сохранён: %s", parsed_dir)


# ══════════════════════════════════════════════════════════════════════
# Публичный интерфейс
# ══════════════════════════════════════════════════════════════════════

def run_page_parser(
    base_dir: str = "data/books",
    book_name: str | None = None,
    # Шаг 1: извлечение текста
    step1_model: str = "",
    step1_prompt: str = "",
    step1_temperature: float = 0.1,
    # Шаг 2: BBox
    step2_model: str = "",
    step2_prompt: str = "",
    step2_temperature: float = 0.0,
    # Шаг 3: описание картинок
    step3_model: str = "",
    step3_prompt: str = "",
    step3_temperature: float = 0.1,
    # Общие параметры
    batch_size: int = 10,
    cancel_event: threading.Event | None = None,
    progress_callback: Callable | None = None,
    log_level: str = "INFO",
) -> dict:
    """Запустить AI-парсинг JPG страниц.

    Обходит: base_dir/*/convert_in_jpg/*.jpg  (или base_dir/book_name/...)
    Сохраняет: base_dir/*/parsed/

    Args:
        base_dir:          Корневая папка с книгами.
        book_name:         Имя конкретной книги (папки). None — все книги.
        step1/2/3_model:   Имена моделей (из реестра).
        step1/2/3_prompt:  Тексты промптов.
        step1/2/3_temperature: Температура генерации.
        batch_size:        Сохранять прогресс каждые N файлов.
        cancel_event:      Событие отмены.
        progress_callback: Функция(current, total, current_file, stats).
        log_level:         Уровень логирования.

    Returns:
        dict: total, success, errors, skipped, elapsed_sec, cancelled,
              total_tokens.
    """
    setup_logging(log_level)

    logger.info("=== Запуск AI-парсинга ===")
    logger.info("Базовая папка: %s, книга: %s", base_dir, book_name or "все")

    base = Path(base_dir)
    files = _find_jpg_files(base, book_name=book_name)
    total = len(files)

    if total == 0:
        logger.info("Нет JPG файлов для обработки в %s", base_dir)
        return {
            "total": 0, "success": 0, "errors": 0, "skipped": 0,
            "elapsed_sec": 0.0, "cancelled": False, "total_tokens": 0,
        }

    logger.info("К обработке: %d JPG файлов", total)

    # Группируем файлы по книгам для отдельного сохранения прогресса
    book_files: dict[str, list[tuple[str, Path, Path]]] = {}
    for book_name, jpg_path, parsed_dir in files:
        book_files.setdefault(book_name, []).append((book_name, jpg_path, parsed_dir))

    t0 = time.perf_counter()
    success = 0
    errors = 0
    skipped = 0
    cancelled = False
    grand_total_tokens = 0
    processed_in_run = 0

    for book_name, book_file_list in book_files.items():
        # Определяем parsed_dir (один на книгу)
        _parsed_dir = book_file_list[0][2]

        # Загружаем предыдущий прогресс для этой книги
        csv_rows, images_data, processed_set = _load_previous_progress(_parsed_dir)

        for jpg_path_item in book_file_list:
            _book, jpg_path, parsed_dir = jpg_path_item
            filename = jpg_path.name

            # Проверка отмены
            if cancel_event and cancel_event.is_set():
                logger.info("Парсинг отменён на файле %s", filename)
                cancelled = True
                break

            # Пропуск уже обработанных
            if filename in processed_set:
                logger.info("⏩ Пропуск: %s (уже обработан)", filename)
                skipped += 1
                processed_in_run += 1
                continue

            relative_name = f"{book_name}/convert_in_jpg/{filename}"
            logger.info("[%d/%d] %s", processed_in_run + 1, total, relative_name)

            try:
                # Читаем изображение
                image_bytes = jpg_path.read_bytes()

                result = process_single_page(
                    image_bytes=image_bytes,
                    filename=filename,
                    step1_model=step1_model,
                    step1_prompt=step1_prompt,
                    step1_temperature=step1_temperature,
                    step2_model=step2_model,
                    step2_prompt=step2_prompt,
                    step2_temperature=step2_temperature,
                    step3_model=step3_model,
                    step3_prompt=step3_prompt,
                    step3_temperature=step3_temperature,
                )

                ts = result["token_stats"]
                grand_total_tokens += ts.get("grand_total", 0)

                # ── Сохранение результатов ──────────────────────────

                # 1. Markdown
                parsed_dir.mkdir(parents=True, exist_ok=True)
                md_name = f"{jpg_path.stem}.md"
                md_path = parsed_dir / md_name
                md_path.write_text(result["text_content"], encoding="utf-8")
                logger.info("Сохранён: %s/%s", book_name, md_name)

                # 2. Вырезанные картинки
                images_dir = parsed_dir / "images"
                for crop_buf, crop_name in zip(
                    result["cropped_images"], result["cropped_names"]
                ):
                    images_dir.mkdir(parents=True, exist_ok=True)
                    (images_dir / crop_name).write_bytes(crop_buf.getvalue())
                    logger.info("Сохранена картинка: %s/%s", book_name, crop_name)

                # 3. Метаданные изображений
                images_data.extend(result["image_data"])

                # 4. CSV-строка
                csv_rows.append([
                    filename, result["duration"], result["has_images"],
                    ts["s1_prompt"], ts["s1_comp"],
                    ts["s2_prompt"], ts["s2_comp"],
                    ts["s3_prompt"], ts["s3_comp"],
                    ts.get("total_prompt", 0), ts.get("total_comp", 0),
                    ts.get("grand_total", 0),
                ])

                processed_set.add(filename)
                success += 1
                processed_in_run += 1

                # Батчевое сохранение
                if success % batch_size == 0:
                    _save_progress(parsed_dir, csv_rows, images_data)

            except FatalAPIError as e:
                errors += 1
                processed_in_run += 1
                logger.error(
                    "❌ ФАТАЛЬНАЯ ОШИБКА API при обработке %s: %s. "
                    "Парсинг остановлен — пополните баланс или проверьте ключ.",
                    filename, e,
                )
                # Сохраняем прогресс и выходим
                if csv_rows:
                    _save_progress(_parsed_dir, csv_rows, images_data)
                cancelled = True
                break

            except Exception as e:
                errors += 1
                processed_in_run += 1
                logger.exception("❌ Ошибка при обработке %s: %s", filename, e)

            # Обратный вызов прогресса
            if progress_callback:
                stats = {
                    "success": success,
                    "errors": errors,
                    "skipped": skipped,
                    "total_tokens": grand_total_tokens,
                }
                progress_callback(processed_in_run, total, relative_name, stats)

        # Сохраняем прогресс книги после завершения
        if csv_rows:
            _save_progress(_parsed_dir, csv_rows, images_data)

        if cancelled:
            break

    elapsed = round(time.perf_counter() - t0, 2)
    logger.info(
        "=== Парсинг завершён: %d успешно, %d ошибок, %d пропущено%s, "
        "%.2fс, %d токенов ===",
        success, errors, skipped,
        ", ОТМЕНЕНО" if cancelled else "",
        elapsed, grand_total_tokens,
    )

    return {
        "total": total,
        "success": success,
        "errors": errors,
        "skipped": skipped,
        "elapsed_sec": elapsed,
        "cancelled": cancelled,
        "total_tokens": grand_total_tokens,
    }


# ══════════════════════════════════════════════════════════════════════
# XLSX-отчёт парсинга (общий с дозаписью)
# ══════════════════════════════════════════════════════════════════════

XLSX_COLUMNS = [
    "timestamp",
    "book_name",
    "filename",
    "pipeline_step",       # "full", "step1", "step2", "step3"
    "duration_sec",
    "has_images",
    "images_count",
    "model",
    "prompt_tokens",
    "completion_tokens",
    "total_tokens",
]

DEFAULT_XLSX_PATH = DATA_DIR / "parsing_report.xlsx"


def _ensure_xlsx(filepath: Path) -> None:
    """Создать xlsx с заголовками, если не существует."""
    import openpyxl

    if filepath.exists():
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Парсинг"
    ws.append(XLSX_COLUMNS)

    widths = {
        "timestamp": 20, "book_name": 30, "filename": 20,
        "pipeline_step": 14, "duration_sec": 12, "has_images": 12,
        "images_count": 14, "model": 25,
        "prompt_tokens": 14, "completion_tokens": 18, "total_tokens": 14,
    }
    for idx, col in enumerate(XLSX_COLUMNS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = widths.get(col, 15)

    wb.save(filepath)


def _append_xlsx_row(
    filepath: Path,
    book_name: str,
    filename: str,
    pipeline_step: str,
    duration: float,
    has_images: bool,
    images_count: int,
    model: str,
    prompt_tokens: int,
    completion_tokens: int,
    total_tokens: int,
) -> None:
    """Добавить одну строку в XLSX-отчёт."""
    import openpyxl
    from datetime import datetime

    _ensure_xlsx(filepath)

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        book_name,
        filename,
        pipeline_step,
        round(duration, 2),
        has_images,
        images_count,
        model,
        prompt_tokens,
        completion_tokens,
        total_tokens,
    ])
    wb.save(filepath)


def _batch_append_xlsx(filepath: Path, rows: list[dict]) -> None:
    """Пакетная дозапись строк в XLSX."""
    import openpyxl
    from datetime import datetime

    if not rows:
        return

    _ensure_xlsx(filepath)
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    for r in rows:
        ws.append([
            r.get("timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            r.get("book_name", ""),
            r.get("filename", ""),
            r.get("pipeline_step", ""),
            r.get("duration_sec", 0),
            r.get("has_images", False),
            r.get("images_count", 0),
            r.get("model", ""),
            r.get("prompt_tokens", 0),
            r.get("completion_tokens", 0),
            r.get("total_tokens", 0),
        ])

    wb.save(filepath)


def _load_xlsx_processed(filepath: Path, pipeline_step: str) -> set[str]:
    """Прочитать XLSX и вернуть множество composite keys 'book_name|filename' для данного шага."""
    import openpyxl

    if not filepath.exists():
        return set()

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    processed = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 4:
            continue
        step_val = row[3]  # pipeline_step
        fname = row[2]     # filename
        book = row[1]      # book_name
        if step_val == pipeline_step and fname:
            processed.add(f"{book}|{fname}")

    wb.close()
    return processed


# ══════════════════════════════════════════════════════════════════════
# Вспомогательные функции для шагов
# ══════════════════════════════════════════════════════════════════════

def _find_parsed_md_files(
    base_dir: Path,
    book_name: str | None = None,
) -> list[tuple[str, Path, Path]]:
    """Найти все .md в parsed/ (результаты шага 1).

    Returns:
        Список: (book_name, md_path, jpg_path)
    """
    results: list[tuple[str, Path, Path]] = []

    if not base_dir.exists():
        return results

    if book_name:
        book_dirs = [base_dir / book_name]
    else:
        book_dirs = sorted(base_dir.iterdir())

    for book_dir in book_dirs:
        if not book_dir.is_dir():
            continue

        parsed_dir = book_dir / "parsed"
        if not parsed_dir.is_dir():
            continue

        jpg_src_dir = book_dir / "convert_in_jpg"

        for md_file in sorted(parsed_dir.glob("*.md")):
            # Ищем соответствующий JPG
            stem = md_file.stem
            jpg_path = None
            for ext in JPG_EXTS:
                candidate = jpg_src_dir / f"{stem}{ext}"
                if candidate.exists():
                    jpg_path = candidate
                    break

            if jpg_path:
                results.append((book_dir.name, md_file, jpg_path))

    return results


def _find_pages_with_crops(
    base_dir: Path,
    book_name: str | None = None,
) -> list[tuple[str, Path, Path, Path]]:
    """Найти страницы с вырезанными картинками и BBox (результаты шага 2).

    Returns:
        Список: (book_name, md_path, bboxes_json_path, images_dir)
    """
    results: list[tuple[str, Path, Path, Path]] = []

    if not base_dir.exists():
        return results

    if book_name:
        book_dirs = [base_dir / book_name]
    else:
        book_dirs = sorted(base_dir.iterdir())

    for book_dir in book_dirs:
        if not book_dir.is_dir():
            continue

        parsed_dir = book_dir / "parsed"
        bboxes_dir = parsed_dir / "bboxes"
        images_dir = parsed_dir / "images"

        if not parsed_dir.is_dir():
            continue

        for md_file in sorted(parsed_dir.glob("*.md")):
            stem = md_file.stem
            bboxes_json = bboxes_dir / f"{stem}.json"

            # Нужен BBox JSON — это значит шаг 2 был выполнен
            if not bboxes_json.exists():
                continue

            results.append((book_dir.name, md_file, bboxes_json, images_dir))

    return results


# ══════════════════════════════════════════════════════════════════════
# Оркестраторы отдельных шагов
# ══════════════════════════════════════════════════════════════════════

def run_step1_only(
    base_dir: str = "data/books",
    book_name: str | None = None,
    model_name: str = "",
    prompt_text: str = "",
    temperature: float = 0.1,
    batch_size: int = 10,
    cancel_event: threading.Event | None = None,
    progress_callback: Callable | None = None,
    log_level: str = "INFO",
    xlsx_path: str | None = None,
) -> dict:
    """Запустить только шаг 1: извлечение текста.

    **Вход:** ``base_dir/*/convert_in_jpg/*.jpg``
    **Выход:** ``base_dir/*/parsed/{stem}.md`` — Markdown с тегами ``[IMAGE_N]``
    **Отчёт:** дозапись в ``data/parsing_report.xlsx`` (pipeline_step = "step1")

    Returns:
        dict: total, success, errors, skipped, elapsed_sec, cancelled, total_tokens
    """
    setup_logging(log_level)

    logger.info("=== Запуск Шага 1: извлечение текста ===")
    logger.info("Модель: %s", model_name)

    xlsx = Path(xlsx_path) if xlsx_path else DEFAULT_XLSX_PATH

    base = Path(base_dir)
    files = _find_jpg_files(base, book_name=book_name)
    total = len(files)

    if total == 0:
        return {"total": 0, "success": 0, "errors": 0, "skipped": 0,
                "elapsed_sec": 0.0, "cancelled": False, "total_tokens": 0}

    # Загружаем уже обработанные из XLSX
    processed_set = _load_xlsx_processed(xlsx, "step1")

    t0 = time.perf_counter()
    success = 0
    errors = 0
    skipped = 0
    cancelled = False
    grand_total_tokens = 0
    xlsx_rows: list[dict] = []

    for idx, (bk, jpg_path, parsed_dir) in enumerate(files):
        filename = jpg_path.name

        if cancel_event and cancel_event.is_set():
            cancelled = True
            break

        if f"{bk}|{filename}" in processed_set:
            logger.info("⏩ Шаг 1: пропуск %s (уже обработан)", filename)
            skipped += 1
            continue

        relative_name = f"{bk}/convert_in_jpg/{filename}"
        logger.info("[Шаг 1] [%d/%d] %s", idx + 1, total, relative_name)

        try:
            image_bytes = jpg_path.read_bytes()

            result = process_page_step1(
                image_bytes=image_bytes,
                filename=filename,
                model_name=model_name,
                prompt_text=prompt_text,
                temperature=temperature,
            )

            grand_total_tokens += result["total_tokens"]

            # Сохраняем .md с тегами
            parsed_dir.mkdir(parents=True, exist_ok=True)
            md_path = parsed_dir / f"{jpg_path.stem}.md"
            md_path.write_text(result["text_content"], encoding="utf-8")

            xlsx_rows.append({
                "book_name": bk,
                "filename": filename,
                "pipeline_step": "step1",
                "duration_sec": result["duration"],
                "has_images": result["has_images"],
                "images_count": result["images_count"],
                "model": model_name,
                "prompt_tokens": result["prompt_tokens"],
                "completion_tokens": result["completion_tokens"],
                "total_tokens": result["total_tokens"],
            })

            processed_set.add(f"{bk}|{filename}")
            success += 1

            if success % batch_size == 0 and xlsx_rows:
                _batch_append_xlsx(xlsx, xlsx_rows)
                xlsx_rows.clear()

        except FatalAPIError as e:
            errors += 1
            logger.error("❌ ФАТАЛЬНАЯ ОШИБКА: %s", e)
            if xlsx_rows:
                _batch_append_xlsx(xlsx, xlsx_rows)
                xlsx_rows.clear()
            cancelled = True
            break

        except Exception as e:
            errors += 1
            logger.exception("❌ Ошибка шага 1 для %s: %s", filename, e)

        if progress_callback:
            stats = {"success": success, "errors": errors, "skipped": skipped,
                     "total_tokens": grand_total_tokens}
            progress_callback(idx + 1, total, relative_name, stats)

    # Финальное сохранение
    if xlsx_rows:
        _batch_append_xlsx(xlsx, xlsx_rows)

    elapsed = round(time.perf_counter() - t0, 2)
    return {"total": total, "success": success, "errors": errors,
            "skipped": skipped, "elapsed_sec": elapsed,
            "cancelled": cancelled, "total_tokens": grand_total_tokens}


def run_step2_only(
    base_dir: str = "data/books",
    book_name: str | None = None,
    model_name: str = "",
    prompt_text: str = "",
    temperature: float = 0.0,
    batch_size: int = 10,
    cancel_event: threading.Event | None = None,
    progress_callback: Callable | None = None,
    log_level: str = "INFO",
    xlsx_path: str | None = None,
) -> dict:
    """Запустить только шаг 2: BBox + вырезка изображений.

    **Вход (от шага 1):**
      - ``base_dir/*/convert_in_jpg/*.jpg`` — исходные изображения
      - ``base_dir/*/parsed/{stem}.md`` — Markdown с тегами ``[IMAGE_N]``
    **Выход:**
      - ``base_dir/*/parsed/bboxes/{stem}.json`` — BBox координаты
      - ``base_dir/*/parsed/images/{stem}_image_N.jpg`` — вырезанные картинки
    **Отчёт:** дозапись в ``data/parsing_report.xlsx`` (pipeline_step = "step2")

    Returns:
        dict: total, success, errors, skipped, elapsed_sec, cancelled, total_tokens
    """
    setup_logging(log_level)

    logger.info("=== Запуск Шага 2: BBox + вырезка ===")
    logger.info("Модель: %s", model_name)

    xlsx = Path(xlsx_path) if xlsx_path else DEFAULT_XLSX_PATH

    base = Path(base_dir)
    md_files = _find_parsed_md_files(base, book_name=book_name)
    total = len(md_files)

    if total == 0:
        logger.warning("Нет .md файлов в parsed/. Сначала запустите шаг 1.")
        return {"total": 0, "success": 0, "errors": 0, "skipped": 0,
                "elapsed_sec": 0.0, "cancelled": False, "total_tokens": 0}

    processed_set = _load_xlsx_processed(xlsx, "step2")

    t0 = time.perf_counter()
    success = 0
    errors = 0
    skipped = 0
    cancelled = False
    grand_total_tokens = 0
    xlsx_rows: list[dict] = []

    for idx, (bk, md_path, jpg_path) in enumerate(md_files):
        filename = jpg_path.name

        if cancel_event and cancel_event.is_set():
            cancelled = True
            break

        if f"{bk}|{filename}" in processed_set:
            logger.info("⏩ Шаг 2: пропуск %s (уже обработан)", filename)
            skipped += 1
            continue

        relative_name = f"{bk}/parsed/{md_path.name}"
        logger.info("[Шаг 2] [%d/%d] %s", idx + 1, total, relative_name)

        try:
            # Читаем Markdown от шага 1 (с тегами)
            markdown_text = md_path.read_text(encoding="utf-8")

            # Если нет IMAGE-тегов — пропускаем
            import re
            if not re.findall(r'\[IMAGE_(\d+)\]', markdown_text):
                logger.info("Шаг 2: пропуск %s — нет IMAGE-тегов", filename)
                # Всё равно записываем в XLSX как обработанный
                xlsx_rows.append({
                    "book_name": bk, "filename": filename,
                    "pipeline_step": "step2", "duration_sec": 0,
                    "has_images": False, "images_count": 0,
                    "model": model_name,
                    "prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0,
                })
                processed_set.add(f"{bk}|{filename}")
                skipped += 1
                continue

            image_bytes = jpg_path.read_bytes()

            result = process_page_step2(
                image_bytes=image_bytes,
                filename=filename,
                markdown_text=markdown_text,
                model_name=model_name,
                prompt_text=prompt_text,
                temperature=temperature,
            )

            grand_total_tokens += result["total_tokens"]

            parsed_dir = md_path.parent
            bboxes_dir = parsed_dir / "bboxes"
            images_dir = parsed_dir / "images"

            # Сохраняем BBox JSON
            bboxes_dir.mkdir(parents=True, exist_ok=True)
            bboxes_json_path = bboxes_dir / f"{jpg_path.stem}.json"
            with open(bboxes_json_path, "w", encoding="utf-8") as f:
                json.dump(result["bboxes"], f, ensure_ascii=False, indent=2)

            # Сохраняем вырезанные картинки
            for crop_buf, crop_name in zip(result["cropped_images"], result["cropped_names"]):
                images_dir.mkdir(parents=True, exist_ok=True)
                (images_dir / crop_name).write_bytes(crop_buf.getvalue())

            xlsx_rows.append({
                "book_name": bk,
                "filename": filename,
                "pipeline_step": "step2",
                "duration_sec": result["duration"],
                "has_images": not result["skipped"],
                "images_count": len(result["cropped_images"]),
                "model": model_name,
                "prompt_tokens": result["prompt_tokens"],
                "completion_tokens": result["completion_tokens"],
                "total_tokens": result["total_tokens"],
            })

            processed_set.add(f"{bk}|{filename}")
            success += 1

            if success % batch_size == 0 and xlsx_rows:
                _batch_append_xlsx(xlsx, xlsx_rows)
                xlsx_rows.clear()

        except FatalAPIError as e:
            errors += 1
            logger.error("❌ ФАТАЛЬНАЯ ОШИБКА: %s", e)
            if xlsx_rows:
                _batch_append_xlsx(xlsx, xlsx_rows)
                xlsx_rows.clear()
            cancelled = True
            break

        except Exception as e:
            errors += 1
            logger.exception("❌ Ошибка шага 2 для %s: %s", filename, e)

        if progress_callback:
            stats = {"success": success, "errors": errors, "skipped": skipped,
                     "total_tokens": grand_total_tokens}
            progress_callback(idx + 1, total, relative_name, stats)

    if xlsx_rows:
        _batch_append_xlsx(xlsx, xlsx_rows)

    elapsed = round(time.perf_counter() - t0, 2)
    return {"total": total, "success": success, "errors": errors,
            "skipped": skipped, "elapsed_sec": elapsed,
            "cancelled": cancelled, "total_tokens": grand_total_tokens}


def run_step3_only(
    base_dir: str = "data/books",
    book_name: str | None = None,
    model_name: str = "",
    prompt_text: str = "",
    temperature: float = 0.1,
    batch_size: int = 10,
    cancel_event: threading.Event | None = None,
    progress_callback: Callable | None = None,
    log_level: str = "INFO",
    xlsx_path: str | None = None,
) -> dict:
    """Запустить только шаг 3: описание изображений + очистка тегов.

    **Вход (от шагов 1 и 2):**
      - ``base_dir/*/parsed/{stem}.md`` — Markdown с тегами ``[IMAGE_N]`` (шаг 1)
      - ``base_dir/*/parsed/bboxes/{stem}.json`` — BBox координаты (шаг 2)
      - ``base_dir/*/parsed/images/{stem}_image_N.jpg`` — вырезанные картинки (шаг 2)
    **Выход:**
      - ``base_dir/*/parsed/{stem}.md`` — перезаписан, теги удалены
      - ``base_dir/*/parsed/images_descriptions.json`` — пополняется
    **Отчёт:** дозапись в ``data/parsing_report.xlsx`` (pipeline_step = "step3")

    Returns:
        dict: total, success, errors, skipped, elapsed_sec, cancelled, total_tokens
    """
    setup_logging(log_level)

    logger.info("=== Запуск Шага 3: описание изображений ===")
    logger.info("Модель: %s", model_name)

    xlsx = Path(xlsx_path) if xlsx_path else DEFAULT_XLSX_PATH

    base = Path(base_dir)
    crop_pages = _find_pages_with_crops(base, book_name=book_name)
    total = len(crop_pages)

    if total == 0:
        logger.warning("Нет страниц с вырезанными картинками. Сначала запустите шаги 1 и 2.")
        return {"total": 0, "success": 0, "errors": 0, "skipped": 0,
                "elapsed_sec": 0.0, "cancelled": False, "total_tokens": 0}

    processed_set = _load_xlsx_processed(xlsx, "step3")

    # Группируем по книгам для images_descriptions.json
    from collections import defaultdict
    book_all_pages: dict[str, list[tuple[str, Path, Path, Path]]] = defaultdict(list)
    for bk, md_path, bboxes_json, images_dir in crop_pages:
        book_all_pages[bk].append((bk, md_path, bboxes_json, images_dir))

    t0 = time.perf_counter()
    success = 0
    errors = 0
    skipped = 0
    cancelled = False
    grand_total_tokens = 0
    xlsx_rows: list[dict] = []
    processed_in_run = 0

    for bk, pages in book_all_pages.items():
        # Загружаем images_descriptions.json для книги
        parsed_dir = pages[0][1].parent
        desc_json_path = parsed_dir / "images_descriptions.json"
        all_image_data: list[dict] = []
        if desc_json_path.exists():
            try:
                with open(desc_json_path, "r", encoding="utf-8") as f:
                    all_image_data = json.load(f)
            except Exception as e:
                logger.warning("Ошибка чтения %s: %s", desc_json_path, e)

        # Отслеживаем уже описанные файлы
        described_files = {d["image_filename"] for d in all_image_data if "image_filename" in d}

        for bk2, md_path, bboxes_json, images_dir in pages:
            filename_stem = md_path.stem
            filename = f"{filename_stem}.jpg"

            if cancel_event and cancel_event.is_set():
                cancelled = True
                break

            processed_in_run += 1

            if f"{bk2}|{filename}" in processed_set:
                logger.info("⏩ Шаг 3: пропуск %s (уже обработан)", filename)
                skipped += 1
                continue

            relative_name = f"{bk}/parsed/{md_path.name}"
            logger.info("[Шаг 3] [%d/%d] %s", processed_in_run, total, relative_name)

            try:
                # Читаем Markdown (с тегами)
                markdown_text = md_path.read_text(encoding="utf-8")

                # Читаем BBox JSON
                with open(bboxes_json, "r", encoding="utf-8") as f:
                    bboxes = json.load(f)

                # Читаем вырезанные картинки
                cropped_images: list[BytesIO] = []
                cropped_names: list[str] = []

                import re
                img_numbers = re.findall(r'\[IMAGE_(\d+)\]', markdown_text)
                for num_str in img_numbers:
                    tag_key = f"IMAGE_{num_str}"
                    crop_name = f"{filename_stem}_{tag_key.lower()}.jpg"
                    crop_path = images_dir / crop_name

                    if crop_path.exists():
                        cropped_images.append(BytesIO(crop_path.read_bytes()))
                        cropped_names.append(crop_name)

                if not cropped_images:
                    logger.info("Шаг 3: пропуск %s — нет вырезанных картинок", filename)
                    xlsx_rows.append({
                        "book_name": bk, "filename": filename,
                        "pipeline_step": "step3", "duration_sec": 0,
                        "has_images": False, "images_count": 0,
                        "model": model_name,
                        "prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0,
                    })
                    processed_set.add(f"{bk2}|{filename}")
                    skipped += 1
                    continue

                result = process_page_step3(
                    markdown_text=markdown_text,
                    filename=filename,
                    cropped_images=cropped_images,
                    cropped_names=cropped_names,
                    bboxes=bboxes,
                    model_name=model_name,
                    prompt_text=prompt_text,
                    temperature=temperature,
                )

                grand_total_tokens += result["total_tokens"]

                # Перезаписываем .md без тегов
                md_path.write_text(result["text_content"], encoding="utf-8")

                # Добавляем описания
                all_image_data.extend(result["image_data"])

                xlsx_rows.append({
                    "book_name": bk,
                    "filename": filename,
                    "pipeline_step": "step3",
                    "duration_sec": result["duration"],
                    "has_images": True,
                    "images_count": len(result["image_data"]),
                    "model": model_name,
                    "prompt_tokens": result["prompt_tokens"],
                    "completion_tokens": result["completion_tokens"],
                    "total_tokens": result["total_tokens"],
                })

                processed_set.add(f"{bk2}|{filename}")
                success += 1

                # Батчевое сохранение JSON + XLSX
                if success % batch_size == 0:
                    with open(desc_json_path, "w", encoding="utf-8") as f:
                        json.dump(all_image_data, f, ensure_ascii=False, indent=2)
                    if xlsx_rows:
                        _batch_append_xlsx(xlsx, xlsx_rows)
                        xlsx_rows.clear()

            except FatalAPIError as e:
                errors += 1
                logger.error("❌ ФАТАЛЬНАЯ ОШИБКА: %s", e)
                # Сохраняем что есть
                with open(desc_json_path, "w", encoding="utf-8") as f:
                    json.dump(all_image_data, f, ensure_ascii=False, indent=2)
                if xlsx_rows:
                    _batch_append_xlsx(xlsx, xlsx_rows)
                    xlsx_rows.clear()
                cancelled = True
                break

            except Exception as e:
                errors += 1
                logger.exception("❌ Ошибка шага 3 для %s: %s", filename, e)

            if progress_callback:
                stats = {"success": success, "errors": errors, "skipped": skipped,
                         "total_tokens": grand_total_tokens}
                progress_callback(processed_in_run, total, relative_name, stats)

        # Сохраняем JSON книги
        with open(desc_json_path, "w", encoding="utf-8") as f:
            json.dump(all_image_data, f, ensure_ascii=False, indent=2)

        if cancelled:
            break

    if xlsx_rows:
        _batch_append_xlsx(xlsx, xlsx_rows)

    elapsed = round(time.perf_counter() - t0, 2)
    return {"total": total, "success": success, "errors": errors,
            "skipped": skipped, "elapsed_sec": elapsed,
            "cancelled": cancelled, "total_tokens": grand_total_tokens}
