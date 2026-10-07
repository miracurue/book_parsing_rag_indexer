"""
Скачивание PDF с arXiv и запись метаданных в XLSX.
Оркестратор — threading + progress_callback + cancel_event.
"""

from __future__ import annotations

import logging
import os
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

import requests
import openpyxl
from openpyxl import Workbook

from src.arxiv_searcher import ArxivArticle, _get_arxiv_user_agent

logger = logging.getLogger(__name__)

XLSX_FILENAME = "articles_metadata.xlsx"
COLUMNS = [
    "arxiv_id",
    "title",
    "authors",
    "published",
    "updated",
    "categories",
    "abstract",
    "pdf_filename",
    "downloaded_at",
]


# ── XLSX-утилиты ──────────────────────────────────────────────

def _ensure_xlsx(xlsx_path: str) -> Workbook:
    """Открывает существующий XLSX или создаёт новый с заголовками."""
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        # Проверяем заголовки
        if ws.cell(1, 1).value != COLUMNS[0]:
            for col, name in enumerate(COLUMNS, 1):
                ws.cell(1, col, name)
            wb.save(xlsx_path)
        return wb

    wb = Workbook()
    ws = wb.active
    ws.title = "Articles"
    for col, name in enumerate(COLUMNS, 1):
        ws.cell(1, col, name)
    wb.save(xlsx_path)
    return wb


def load_downloaded_ids(xlsx_path: str) -> set[str]:
    """Возвращает множество arxiv_id, уже записанных в XLSX."""
    if not os.path.exists(xlsx_path):
        return set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        ws = wb.active
        ids: set[str] = set()
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                ids.add(str(row[0]))
        wb.close()
        return ids
    except Exception:
        return set()


def _append_row(xlsx_path: str, article: ArxivArticle, pdf_filename: str) -> None:
    """Дописывает одну строку в XLSX."""
    wb = _ensure_xlsx(xlsx_path)
    ws = wb.active
    row_num = ws.max_row + 1
    values = [
        article.arxiv_id,
        article.title,
        "; ".join(article.authors),
        article.published,
        article.updated,
        "; ".join(article.categories),
        article.abstract,
        pdf_filename,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    ]
    for col, val in enumerate(values, 1):
        ws.cell(row_num, col, val)
    wb.save(xlsx_path)
    wb.close()


# ── Скачивание одной статьи ────────────────────────────────────

_DOWNLOAD_RETRIES = 3
_DOWNLOAD_BACKOFF = [5.0, 15.0, 30.0]


def download_article_pdf(
    article: ArxivArticle,
    save_dir: str,
) -> Optional[str]:
    """
    Скачивает PDF одной статьи в save_dir через requests.

    Использует requests вместо arxiv.Client.download_pdf(),
    чтобы избежать SSL-ошибок на Windows (urllib не находит CA-bundle).

    Returns
    -------
    str or None
        Имя скачанного файла или None при ошибке.
    """
    save_dir_path = Path(save_dir)
    save_dir_path.mkdir(parents=True, exist_ok=True)
    filename = article.safe_filename()
    filepath = save_dir_path / filename

    # Если файл уже есть — не скачиваем повторно
    if filepath.exists() and filepath.stat().st_size > 0:
        logger.info("PDF уже существует: %s", filepath)
        return filename

    # URL для скачивания PDF напрямую с arXiv
    pdf_url = f"https://arxiv.org/pdf/{article.arxiv_id}.pdf"
    headers = {"User-Agent": _get_arxiv_user_agent()}

    last_exc: Optional[Exception] = None
    for attempt in range(_DOWNLOAD_RETRIES):
        try:
            resp = requests.get(
                pdf_url,
                headers=headers,
                timeout=60,
                allow_redirects=True,
            )
            resp.raise_for_status()

            # Проверяем, что ответ действительно PDF
            content_type = resp.headers.get("Content-Type", "")
            if "pdf" not in content_type and not resp.content[:5] == b"%PDF-":
                logger.warning(
                    "Ответ не PDF для %s (Content-Type: %s, первые байты: %s)",
                    article.arxiv_id, content_type, resp.content[:20],
                )

            filepath.write_bytes(resp.content)
            logger.info("Скачан: %s → %s (%d KB)", article.arxiv_id, filepath, len(resp.content) // 1024)
            return filename

        except requests.exceptions.HTTPError as e:
            last_exc = e
            status = getattr(e.response, "status_code", None)
            delay = _DOWNLOAD_BACKOFF[min(attempt, len(_DOWNLOAD_BACKOFF) - 1)]
            logger.warning(
                "HTTP %s при скачивании %s (попытка %d/%d). Повтор через %.0f сек.",
                status, article.arxiv_id, attempt + 1, _DOWNLOAD_RETRIES, delay,
            )
            if attempt < _DOWNLOAD_RETRIES - 1:
                time.sleep(delay)

        except requests.exceptions.ConnectionError as e:
            last_exc = e
            delay = _DOWNLOAD_BACKOFF[min(attempt, len(_DOWNLOAD_BACKOFF) - 1)]
            logger.warning(
                "Ошибка соединения при скачивании %s (попытка %d/%d): %s. Повтор через %.0f сек.",
                article.arxiv_id, attempt + 1, _DOWNLOAD_RETRIES, e, delay,
            )
            if attempt < _DOWNLOAD_RETRIES - 1:
                time.sleep(delay)

        except Exception as e:
            last_exc = e
            logger.error("Ошибка скачивания %s: %s", article.arxiv_id, e)
            return None

    logger.error("Не удалось скачать %s после %d попыток: %s", article.arxiv_id, _DOWNLOAD_RETRIES, last_exc)
    return None


# ── Оркестратор (для threading) ───────────────────────────────

def run_download_articles(
    articles: list[ArxivArticle],
    save_dir: str,
    xlsx_path: str,
    progress_callback: Optional[Callable[[float, str], None]] = None,
    cancel_event: Optional[threading.Event] = None,
) -> dict:
    """
    Скачивает PDF и записывает метаданные в XLSX.

    Parameters
    ----------
    articles : list[ArxivArticle]
        Статьи для скачивания.
    save_dir : str
        Папка для сохранения PDF.
    xlsx_path : str
        Путь к XLSX-файлу с метаданными.
    progress_callback : callable(progress: float, message: str)
        Функция обратного вызова для прогресса.
    cancel_event : threading.Event
        Событие отмены.

    Returns
    -------
    dict with keys: total, downloaded, skipped, errors, error_details
    """
    cancel_event = cancel_event or threading.Event()
    total = len(articles)
    if total == 0:
        return {"total": 0, "downloaded": 0, "skipped": 0, "errors": 0, "error_details": []}

    # Дедупликация по XLSX
    already_ids = load_downloaded_ids(xlsx_path)
    os.makedirs(save_dir, exist_ok=True)

    downloaded = 0
    skipped = 0
    errors = 0
    error_details: list[str] = []

    for i, article in enumerate(articles):
        if cancel_event.is_set():
            logger.info("Скачивание отменено на статье %d/%d", i, total)
            break

        msg = f"({i + 1}/{total}) {article.arxiv_id}: {article.title[:60]}..."
        if progress_callback:
            progress_callback((i) / total, msg)

        # Пропускаем уже скачанные
        if article.arxiv_id in already_ids:
            # Но проверяем, есть ли PDF на диске
            pdf_path = Path(save_dir) / article.safe_filename()
            if pdf_path.exists():
                skipped += 1
                logger.info("Пропуск (уже в XLSX + PDF): %s", article.arxiv_id)
                continue

        # Скачиваем PDF
        filename = download_article_pdf(article, save_dir)
        if filename is None:
            errors += 1
            error_details.append(f"{article.arxiv_id}: ошибка скачивания PDF")
            continue

        # Записываем в XLSX
        try:
            _append_row(xlsx_path, article, filename)
            downloaded += 1
        except Exception as e:
            errors += 1
            error_details.append(f"{article.arxiv_id}: ошибка записи XLSX — {e}")

        # Задержка между запросами (arXiv: ≤1 запрос / 3 сек)
        if i < total - 1 and not cancel_event.is_set():
            time.sleep(3)

    if progress_callback:
        progress_callback(1.0, f"Готово! Скачано: {downloaded}, пропущено: {skipped}, ошибок: {errors}")

    return {
        "total": total,
        "downloaded": downloaded,
        "skipped": skipped,
        "errors": errors,
        "error_details": error_details,
    }


# ── Утилиты для UI ────────────────────────────────────────────

def get_existing_subfolders(base_dir: str) -> list[str]:
    """Возвращает список подпапок в base_dir."""
    if not os.path.isdir(base_dir):
        return []
    return sorted([
        d for d in os.listdir(base_dir)
        if os.path.isdir(os.path.join(base_dir, d))
    ])


def get_subfolder_stats(base_dir: str, subfolder: str) -> dict:
    """Статистика по подпапке: кол-во PDF, есть ли XLSX."""
    sf_path = os.path.join(base_dir, subfolder)
    pdf_count = len(list(Path(sf_path).glob("*.pdf")))
    xlsx_exists = os.path.exists(os.path.join(sf_path, XLSX_FILENAME))
    xlsx_rows = 0
    if xlsx_exists:
        ids = load_downloaded_ids(os.path.join(sf_path, XLSX_FILENAME))
        xlsx_rows = len(ids)
    return {"pdf_count": pdf_count, "xlsx_exists": xlsx_exists, "xlsx_rows": xlsx_rows}